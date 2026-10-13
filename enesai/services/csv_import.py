"""Product import from Excel (.xlsx) or CSV (V1-13, ADMIN_REQUIREMENTS §5).

Upload → parse (.xlsx first sheet via openpyxl; CSV UTF-8 with/without BOM, fallback Windows-1252,
delimiter , ; or TAB; the market export “Kodi · Përshkrimi · Njësi” is recognised) →
per-row validation and preview (update / new / warning / error) → run in batches of 500
inside transactions in a background thread with progress → results + error CSV.
Rows match existing products by SKU (update) else create. Nothing is ever deleted.
foto_url is NOT downloaded (SSRF risk); only bundled /static/img/… paths are accepted."""
import csv
import io
import json
import re
import threading
from pathlib import Path

from flask import current_app

from ..db import connect, get_db
from ..util import csv_safe, norm, now_iso, slugify
from . import audit
from . import products as psvc
from .catalog import DISCOUNTS

COLUMNS = ["sku", "emri_sq", "emri_en", "kategoria", "nenkategoria", "marka", "paketimi", "njesia", "barkodi", "cmimi", "zbritja", "stoku",
           "i_dukshem", "foto_url"]
# Header names seen in shop/market exports (accents and case are ignored; spaces → “_”).
ALIASES = {"emri": "emri_sq", "name": "emri_sq", "name_sq": "emri_sq", "name_en": "emri_en", "category": "kategoria", "subcategory": "nenkategoria",
           "brand": "marka", "pack": "paketimi", "packaging": "paketimi", "price": "cmimi", "discount": "zbritja", "stock": "stoku",
           "visible": "i_dukshem", "image": "foto_url", "photo": "foto_url",
           # market program export (ARTIKUJ_TOTAL.xlsx: Kodi · Përshkrimi · Njësi) and common variants
           "kodi": "sku", "kodi_i_artikullit": "sku", "kod": "sku", "code": "sku",
           "pershkrimi": "emri_sq", "artikulli": "emri_sq", "emertimi": "emri_sq", "emri_i_artikullit": "emri_sq",
           "njesi": "njesia", "njesia_matese": "njesia", "nj": "njesia", "unit": "njesia",
           "barkod": "barkodi", "barcode": "barkodi", "ean": "barkodi",
           "cmim": "cmimi", "cmimi_shitjes": "cmimi", "cmimi_i_shitjes": "cmimi", "cmimi_me_tvsh": "cmimi", "cmim_shitje": "cmimi",
           "sasia": "stoku", "gjendja": "stoku", "stok": "stoku", "sasi": "stoku",
           "grupi": "kategoria", "kategori": "kategoria", "nengrupi": "nenkategoria"}
UNIT_MAP = {"cope": "copë", "cp": "copë", "copa": "copë", "pcs": "copë", "kg": "kg", "kilogram": "kg", "l": "L", "lt": "L", "liter": "L",
            "litra": "L", "m": "m", "metra": "m", "meter": "m", "pako": "pako", "pk": "pako"}
MAX_ROWS = 200_000
BATCH = 500


class ImportError_(Exception):
    pass


def decode(raw):
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def sniff_delimiter(text):
    first = text.split("\n", 1)[0]
    counts = {d: first.count(d) for d in (",", ";", "\t")}
    d = max(counts, key=counts.get)
    return d if counts[d] else ","


def parse_price(v):
    s = str(v or "").strip().replace(" ", "").replace("L", "").replace("l", "")
    if re.match(r"^\d{1,3}(\.\d{3})+$", s):
        s = s.replace(".", "")
    m = re.match(r"^(\d+)(?:[.,]0+)?$", s)
    if not m:
        return None
    n = int(m.group(1))
    return n if 0 < n <= 10_000_000 else None


def _cell(v):
    """Excel cell → text. Whole numbers lose the “.0” (codes, prices, stock)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, int):
        return str(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v).strip()


def read_table(raw, filename=""):
    """→ (header, iterator of (line_number, cells)). Excel .xlsx (first sheet) or CSV/TXT."""
    name = (filename or "").lower()
    if name.endswith(".xls") or raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ImportError_("Formati i vjetër .xls nuk pranohet. Hapeni në Excel dhe ruajeni si .xlsx (Excel Workbook) ose CSV.")
    if name.endswith(".xlsx") or raw[:4] == b"PK\x03\x04":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise ImportError_("Skedari Excel nuk u lexua. Ruajeni sërish si .xlsx dhe provoni përsëri.")
        ws = next((w for w in wb.worksheets if (w.max_row or 0) > 0), None)
        if ws is None:
            raise ImportError_("Skedari është bosh.")
        it = ((i, [_cell(c) for c in row]) for i, row in enumerate(ws.iter_rows(values_only=True), start=1))
    else:
        text = decode(raw)
        if not text.strip():
            raise ImportError_("Skedari është bosh.")
        rdr = csv.reader(io.StringIO(text), delimiter=sniff_delimiter(text))
        it = ((i, [c.strip() for c in row]) for i, row in enumerate(rdr, start=1))
    for i, cells in it:
        if any(cells):
            return cells, it
    raise ImportError_("Skedari është bosh.")


def _col(h):
    k = norm(h).replace(" ", "_").replace(".", "")
    return ALIASES.get(k, k)


def clean_sku(v):
    """Market codes sometimes contain spaces (“DJATH GOUDA”, “6935031405606 .”): trim stray
    dots/spaces and join words with “-”, the same way on every import so updates still match."""
    return re.sub(r"\s+", "-", str(v or "").strip().strip(".").strip())


def parse(raw, auto_categories=True, con=None, filename="", backorder=False):
    """→ (rows, notes). Only sku + name are required; price, category and stock may be missing
    (products are then stored hidden and completed by a later import or in the admin)."""
    con = con or get_db()
    header, it = read_table(raw, filename)
    cols = [_col(h) for h in header]
    if "sku" not in cols:
        raise ImportError_("Mungon kolona e kodit (SKU / Kodi). Duhen të paktën kodi dhe — për produktet e reja — emri (Emri / Përshkrimi).")
    has = set(cols)
    notes = []
    if "cmimi" not in has:
        notes.append("Skedari nuk ka kolonë çmimi: produktet e reja ruhen të fshehura (pa çmim). Kur të keni skedarin me çmime, "
                     "importojeni sërish — produktet përditësohen sipas kodit dhe dalin në dyqan sapo të kenë çmim dhe kategori.")
    if "kategoria" not in has:
        notes.append("Skedari nuk ka kolonë kategorie: produktet e reja mbeten “Pa kategori” (të fshehura). Caktojuani kategorinë te "
                     "Produktet › filtri “Pa kategori” (kërkoni p.sh. “AKULLORE” dhe zbatoni kategorinë për të gjitha).")
    if "stoku" not in has:
        notes.append("Skedari nuk ka stok: stoku i produkteve ekzistuese nuk ndryshon; produktet e reja fillojnë me 0"
                     + (" dhe lejohen porositë edhe pa stok (zgjodhët “stoku nuk ndiqet”)." if backorder
                        else " (“Mbaroi”). Nëse nuk e ndiqni stokun online, zgjidhni opsionin “stoku nuk ndiqet” në ngarkim."))
    cats = {}
    for r in con.execute("SELECT id, slug, name_sq, parent_id FROM categories"):
        cats[(r["parent_id"], norm(r["name_sq"]))] = r["id"]
        cats[(r["parent_id"], r["slug"])] = r["id"]
    existing = {r[0] for r in con.execute("SELECT sku FROM products")}
    rows, seen = [], set()
    for i, rec in it:
        if len(rows) >= MAX_ROWS:
            raise ImportError_("Shumë rreshta (maksimumi %d)." % MAX_ROWS)
        if not any(rec):
            continue
        d = {}
        for k, c in enumerate(cols):
            if c in COLUMNS and k < len(rec) and c not in d:
                d[c] = rec[k].strip()
        errs, warns = [], []
        raw_sku = d.get("sku", "")
        sku = clean_sku(raw_sku)
        if not sku:
            errs.append(("Mungon SKU / kodi", "Shtoni një kod unik dhe importoni sërish."))
        elif not psvc.SKU_RE.match(sku):
            errs.append(("Kodi “%s” ka shenja të palejuara" % raw_sku, "Përdorni shkronja, shifra, “-”, “_” ose “.” (deri 40)."))
        elif sku in seen:
            errs.append(("Kodi “%s” përsëritet në skedar" % sku, "Mbani vetëm një rresht për çdo kod."))
        elif sku != raw_sku:
            warns.append("Kodi u përshtat: “%s” → “%s”" % (raw_sku, sku))
        seen.add(sku)
        if len(d.get("emri_sq", "")) < 2 and sku not in existing:
            errs.append(("Mungon emri i produktit të ri", "Plotësoni emrin / përshkrimin (për kodet ekzistuese mund të mungojë)."))
        price = None
        if d.get("cmimi"):
            price = parse_price(d["cmimi"])
            if price is None:
                errs.append(("Çmimi “%s” nuk është numër" % d["cmimi"], "Përdorni vetëm shifra, p.sh. 250."))
        pct = None
        if d.get("zbritja"):
            try:
                pct = int(d["zbritja"].replace("%", "").strip() or 0)
                if pct not in DISCOUNTS:
                    raise ValueError
            except ValueError:
                errs.append(("Zbritja “%s” nuk lejohet" % d["zbritja"], "Vlerat e lejuara: 0, 5, 10, 15, 20, 25, 30, 40, 50."))
        stock = None
        if d.get("stoku"):
            try:
                st_raw = d["stoku"].replace(" ", "")
                if re.match(r"^\d{1,3}(\.\d{3})+$", st_raw):      # 1.250 → 1250 (thousands)
                    st_raw = st_raw.replace(".", "")
                stock = int(float(st_raw.replace(",", ".")))       # 12,5 kg → 12
                if stock < 0:
                    raise ValueError
            except ValueError:
                errs.append(("Stoku “%s” nuk është numër" % d["stoku"], "Shkruani një numër të plotë ≥ 0."))
        vis = None
        if d.get("i_dukshem"):
            vis = norm(d["i_dukshem"]) not in ("0", "jo", "false", "no", "n")
        cname = d.get("kategoria", "")
        cat_id = cats.get((None, norm(cname))) or cats.get((None, slugify(cname))) if cname else None
        if cname and not cat_id and not auto_categories:
            errs.append(("Kategoria “%s” nuk ekziston" % cname, "Krijojeni te Kategoritë ose aktivizoni krijimin automatik."))
        elif cname and not cat_id:
            warns.append("Kategori e re: %s" % cname)
        sname = d.get("nenkategoria", "")
        if sname and cat_id and not cats.get((cat_id, norm(sname))):
            if auto_categories:
                warns.append("Nënkategori e re: %s" % sname)
            else:
                errs.append(("Nënkategoria “%s” nuk ekziston" % sname, "Krijojeni ose aktivizoni krijimin automatik."))
        unit = ""
        if d.get("njesia"):
            unit = UNIT_MAP.get(norm(d["njesia"]).replace(".", ""), "")
            if not unit:
                warns.append("Njësia “%s” nuk njihet → copë" % d["njesia"])
                unit = "copë"
        barcode = re.sub(r"\s", "", d.get("barkodi", ""))
        if barcode and not barcode.isdigit():
            warns.append("Barkodi “%s” nuk ka vetëm shifra → nuk u ruajt" % barcode)
            barcode = ""
        if not barcode and sku.isdigit() and len(sku) in (8, 12, 13, 14):
            barcode = sku  # the market code is an EAN/UPC barcode
        foto = d.get("foto_url", "")
        if foto and not foto.startswith("/static/img/"):
            warns.append("foto_url nuk shkarkohet — ngarkoni foton te produkti")
            foto = ""
        rows.append({"line": i, "sku": sku, "name_sq": d.get("emri_sq", ""), "name_en": d.get("emri_en", ""), "cat": cname, "sub": sname,
                     "brand": d.get("marka", ""), "pack": d.get("paketimi", ""), "unit": unit, "barcode": barcode, "price": price, "pct": pct,
                     "stock": stock, "visible": vis, "foto": foto, "errors": errs, "warnings": warns,
                     "kind": "error" if errs else ("update" if sku in existing else "new"),
                     "raw_price": d.get("cmimi", ""), "raw_stock": d.get("stoku", "")})
    return rows, notes


def summary(rows):
    return {"total": len(rows), "ok": sum(1 for r in rows if not r["errors"]), "warn": sum(1 for r in rows if r["warnings"] and not r["errors"]),
            "err": sum(1 for r in rows if r["errors"]), "new": sum(1 for r in rows if r["kind"] == "new"), "upd": sum(1 for r in rows if r["kind"] == "update")}


def save_job(filename, rows, auto_categories, operator, notes=(), backorder=False):
    con = get_db()
    s = summary(rows)
    cur = con.execute("INSERT INTO import_jobs(filename, created_by, created_at, status, total, warnings, errors, auto_categories) VALUES (?,?,?, 'preview', ?,?,?,?)",
                      (filename[:120], operator["name"], now_iso(), s["total"], s["warn"], s["err"], 1 if auto_categories else 0))
    jid = cur.lastrowid
    path = Path(current_app.config["PRIVATE_DIR"]) / "imports" / ("job-%d.json" % jid)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"rows": rows, "notes": list(notes), "opts": {"backorder": bool(backorder)}}, ensure_ascii=False), encoding="utf-8")
    con.execute("UPDATE import_jobs SET payload_path=? WHERE id=?", (path.name, jid))
    return jid


def _payload(private_dir, job):
    data = json.loads((Path(private_dir) / "imports" / job["payload_path"]).read_text(encoding="utf-8"))
    return data if isinstance(data, dict) else {"rows": data, "notes": [], "opts": {}}


def load_payload(job):
    return _payload(current_app.config["PRIVATE_DIR"], job)


def load_rows(job):
    return load_payload(job)["rows"]


def start(job_id, operator):
    app = current_app._get_current_object()
    con = get_db()
    cur = con.execute("UPDATE import_jobs SET status='running' WHERE id=? AND status='preview'", (job_id,))
    if cur.rowcount != 1:
        return False
    th = threading.Thread(target=_run, args=(app, job_id, operator), daemon=True)
    th.start()
    if app.config.get("TESTING"):
        th.join()
    return True


def _run(app, job_id, operator):
    with app.app_context():
        con = connect(app.config["DATABASE"])
        try:
            job = con.execute("SELECT * FROM import_jobs WHERE id=?", (job_id,)).fetchone()
            payload = _payload(app.config["PRIVATE_DIR"], job)
            rows, opts = payload["rows"], payload.get("opts") or {}
            good = [r for r in rows if not r["errors"]]
            created = updated = done = 0
            for start_i in range(0, len(good), BATCH):
                chunk = good[start_i:start_i + BATCH]
                con.execute("BEGIN IMMEDIATE")
                try:
                    for r in chunk:
                        c, u = _apply(con, r, job["auto_categories"], opts)
                        created += c; updated += u
                    con.execute("COMMIT")
                except Exception:
                    con.execute("ROLLBACK")
                    raise
                done += len(chunk)
                con.execute("UPDATE import_jobs SET done_rows=?, created_n=?, updated_n=? WHERE id=?", (done, created, updated, job_id))
            errs = [r for r in rows if r["errors"]]
            if errs:
                p = Path(app.config["PRIVATE_DIR"]) / "imports" / ("job-%d-gabimet.csv" % job_id)
                with p.open("w", encoding="utf-8-sig", newline="") as fh:
                    w = csv.writer(fh, delimiter=";")
                    w.writerow(["rreshti", "sku", "emri_sq", "problemi", "si_ta_rregulloni"])
                    for r in errs:
                        for prob, fix in r["errors"]:
                            w.writerow([r["line"], csv_safe(r["sku"]), csv_safe(r["name_sq"]), csv_safe(prob), csv_safe(fix)])
                con.execute("UPDATE import_jobs SET error_csv_path=? WHERE id=?", (p.name, job_id))
            con.execute("UPDATE import_jobs SET status='done', finished_at=? WHERE id=?", (now_iso(), job_id))
            audit.log("Importoi produkte", "%s · %d ok, %d gabime" % (job["filename"], created + updated, len(errs)),
                      "", "krijuar %d · përditësuar %d" % (created, updated), area="import", con=con, operator=operator)
        except Exception:  # pragma: no cover - logged, job marked failed
            app.logger.exception("import job %s failed", job_id)
            con.execute("UPDATE import_jobs SET status='failed', finished_at=? WHERE id=?", (now_iso(), job_id))
        finally:
            con.close()


def _category(con, name, parent_id, auto):
    nm = norm(name)
    for row in con.execute("SELECT id, name_sq, slug FROM categories WHERE parent_id IS ?", (parent_id,)):
        if norm(row["name_sq"]) == nm or row["slug"] == slugify(name):
            return row["id"]
    if not auto:
        return None
    parent_slug = con.execute("SELECT slug FROM categories WHERE id=?", (parent_id,)).fetchone()[0] if parent_id else None
    sort = (con.execute("SELECT COALESCE(MAX(sort),0) FROM categories WHERE parent_id IS ?", (parent_id,)).fetchone()[0] or 0) + 1
    t = now_iso()
    cur = con.execute("INSERT INTO categories(slug, name_sq, parent_id, sort, visible, created_at, updated_at) VALUES (?,?,?,?,1,?,?)",
                      (psvc.unique_slug(name, con, parent_slug=parent_slug), name.strip()[:80], parent_id, sort, t, t))
    return cur.lastrowid


def _apply(con, r, auto, opts=None):
    """Create or update one product. Only the columns present (and non-empty) in the file
    change an existing product; a product without price or category is stored hidden."""
    opts = opts or {}
    ex = con.execute("SELECT * FROM products WHERE sku=?", (r["sku"],)).fetchone()
    g = (lambda k, default: ex[k] if ex else default)
    cat = _category(con, r["cat"], None, auto) if r["cat"] else None
    sub = _category(con, r["sub"], cat, auto) if r["sub"] and cat else None
    cat_id = cat or g("category_id", None)
    sub_id = sub if (sub or r["cat"]) else g("subcategory_id", None)
    price = r["price"] if r["price"] is not None else g("price", 0)
    unit = r.get("unit") or g("unit", "copë")
    pack = r["pack"] or g("packaging", "")
    if not pack and not ex and unit in ("kg", "L", "m"):
        pack = "1 " + unit  # price is per kg / litre / metre
    backorder = g("allow_backorder", 0)
    if opts.get("backorder") and r["stock"] is None:
        backorder = 1
    complete = bool(price and price > 0 and cat_id)
    if r["visible"] is not None:                       # explicit i_dukshem column wins (when possible)
        visible, auto_hidden = (1 if r["visible"] and complete else 0), (1 if r["visible"] and not complete else 0)
    elif not ex:                                         # new: live only when complete
        visible, auto_hidden = (1, 0) if complete else (0, 1)
    elif ex["auto_hidden"] and complete:                 # was waiting for data → goes live now
        visible, auto_hidden = 1, 0
    elif not complete:
        visible, auto_hidden = 0, (1 if ex["visible"] or ex["auto_hidden"] else 0)
    else:
        visible, auto_hidden = ex["visible"], ex["auto_hidden"]
    d = {"sku": r["sku"], "name_sq": (r["name_sq"] or g("name_sq", ""))[:160], "name_en": (r["name_en"] or g("name_en", ""))[:160],
         "desc_sq": g("desc_sq", ""), "desc_en": g("desc_en", ""), "brand": (r["brand"] or g("brand", ""))[:80],
         "packaging": pack[:60], "unit": unit, "barcode": (r.get("barcode") or g("barcode", ""))[:20],
         "price": price, "discount_pct": r["pct"] if r["pct"] is not None else g("discount_pct", 0),
         "stock": r["stock"] if r["stock"] is not None else g("stock", 0), "low_threshold": g("low_threshold", 5),
         "allow_backorder": backorder, "visible": visible, "auto_hidden": auto_hidden, "on_offers": g("on_offers", 1),
         "category_id": cat_id, "subcategory_id": sub_id}
    if ex:
        psvc.update(ex["id"], d, con)
        pid, c, u = ex["id"], 0, 1
    else:
        pid, c, u = psvc.insert(d, con), 1, 0
    if r.get("foto") and not con.execute("SELECT 1 FROM product_images WHERE product_id=?", (pid,)).fetchone():
        con.execute("INSERT INTO product_images(product_id, path, sort) VALUES (?,?,0)", (pid, r["foto"]))
    return c, u


TEMPLATE = "﻿" + ";".join(COLUMNS) + "\n" + "FR-1050;Avokado Hass;Hass avocado;Fruta & Perime;Fruta;Import · Spanjë;2 copë;copë;;300;20;30;1;\n"
